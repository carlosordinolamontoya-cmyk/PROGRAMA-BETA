import os
import tempfile
from docx import Document
import app_kardex as app


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def make_vale(n):
    obj = app.PrintPreviewDialog.__new__(app.PrintPreviewDialog)
    obj.movimiento = {
        "tipo_codigo": "CO",
        "vale": "VALE-2026-000001",
        "fecha_operacion": "2026-08-12 10:00:00",
        "documento": "PRUEBA",
        "ot": "OT-001",
        "centro_costo_codigo": "04",
        "centro_costo_nombre": "LOGISTICA",
        "almacen_codigo": "AQP",
        "solicitante": "USUARIO PRUEBA",
        "usuario": "admin",
        "observacion": "PRUEBA V45",
    }
    obj.detalles = [
        {
            "item": i,
            "codigo": f"ART-{i:03d}",
            "descripcion": f"PRODUCTO {i}",
            "unidad_medida": "UND",
            "cantidad": -i,
        }
        for i in range(1, n + 1)
    ]
    return obj.build_word_document()


class FakeDB:
    def __init__(self, n):
        self.n = n

    def get_orden_compra(self, numero):
        oc = {
            "tipo_orden": "NACIONAL",
            "numero": numero,
            "fecha_orden": "2026-08-12",
            "razon_social": "PROVEEDOR PRUEBA SAC",
            "direccion": "AV. PRUEBA 123",
            "cotizacion": "COT-001",
            "ruc": "20123456789",
            "telefono": "999999999",
            "lugar_entrega": "ALMACEN AREQUIPA",
            "almacen_codigo": "AQP",
            "almacen_nombre": "ALMACEN AREQUIPA",
            "forma_pago": "CONTADO",
            "moneda": "SOLES",
            "subtotal": 1000.0,
            "igv": 180.0,
            "total": 1180.0,
            "usuario_creador": "admin",
            "solicitante": "SOLICITANTE PRUEBA",
            "aprobado_por": "admin",
        }
        det = [
            {
                "item": i,
                "cantidad_aprobada": 0.0,
                "cantidad_solicitada": 1.0,
                "unidad_medida": "UND",
                "descripcion": f"ITEM DE COMPRA NUMERO {i}",
                "precio_unitario_sin_igv": 10.0,
                "subtotal": 10.0,
            }
            for i in range(1, self.n + 1)
        ]
        return oc, det

    def get_proveedor_by_ruc(self, ruc, active_only=False):
        return {"banco": "BCP", "numero_cuenta": "123456789", "cci": "00212345678912345678"}


def make_oc(n):
    obj = app.OCPrintPreviewDialog.__new__(app.OCPrintPreviewDialog)
    obj.numero = "OC-2026-000001"
    obj.db = FakeDB(n)
    return obj.build_word_document()


def item_tables(doc):
    result = []
    for table in doc.tables:
        if not table.rows:
            continue
        hdr = {app.normalize_text(c.text).upper() for c in table.rows[0].cells}
        if {"ITEM", "CÓDIGO", "DESCRIPCIÓN", "UM", "CANTIDAD"}.issubset(hdr):
            result.append(table)
    return result


def test_vale_2_items_compacto():
    doc = make_vale(2)
    tables = item_tables(doc)
    assert len(tables) == 1
    assert len(tables[0].rows) == 3, "Debe quedar cabecera + 2 ítems, sin 10 filas vacías"
    text = all_text(doc)
    assert "01 de 01" in text
    assert "VALE DE SALIDA DE ALMACÉN" in text
    assert "VALE DE SALIDA DE ALMACEN DE ALMACÉN" not in text


def test_vale_13_items_repite_hoja():
    doc = make_vale(13)
    tables = item_tables(doc)
    assert len(tables) == 2
    assert len(tables[0].rows) == 13  # cabecera + 12
    assert len(tables[1].rows) == 2   # cabecera + item 13
    text = all_text(doc)
    assert text.count("SISTEMA INTEGRADO DE GESTIÓN") >= 2
    assert "01 de 02" in text and "02 de 02" in text
    assert "ART-013" in text


def test_vale_25_items_tres_hojas_logicas():
    doc = make_vale(25)
    tables = item_tables(doc)
    assert len(tables) == 3
    assert [len(t.rows) for t in tables] == [13, 13, 2]
    text = all_text(doc)
    assert "01 de 03" in text and "02 de 03" in text and "03 de 03" in text
    assert "ART-025" in text


def test_oc_mas_de_16_agrega_filas_sin_duplicar_documento():
    doc = make_oc(40)
    text = all_text(doc)
    # 6 tablas originales: no se crea un segundo documento/cabecera/cierre.
    assert len(doc.tables) == 6
    assert text.count("SMETAL GROUP S.A.C.") >= 1
    # La razón social aparece una sola vez en el encabezado principal (puede repetirse por celdas combinadas internamente).
    assert "ITEM DE COMPRA NUMERO 40" in text
    assert text.count("INFORMACION PARA EL PROVEEDOR:") >= 1
    assert text.count("ELABORADO POR:") >= 1
    items_table = app._find_docx_table(doc, ["CANT", "UNID", "DESCRIPCION", "VALOR UNIT.", "IMPORTE"])
    assert items_table is not None
    # cabecera + 40 ítems + 3 filas de totales/banco
    assert len(items_table.rows) == 44


def main():
    tests = [
        test_vale_2_items_compacto,
        test_vale_13_items_repite_hoja,
        test_vale_25_items_tres_hojas_logicas,
        test_oc_mas_de_16_agrega_filas_sin_duplicar_documento,
    ]
    for test in tests:
        test()
        print(f"OK {test.__name__}")
    print(f"V45 impresiones: {len(tests)}/{len(tests)} pruebas OK")


if __name__ == "__main__":
    main()
